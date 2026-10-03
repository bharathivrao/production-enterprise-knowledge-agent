import os
from unittest.mock import patch

import pymupdf
import pytest
from docx import Document

from app.db.database import close_database_pool
from app.db.repository import delete_document
from app.ingestion.pipeline import ingest_document


pytestmark = pytest.mark.integration


if os.getenv("RUN_POSTGRES_INTEGRATION") != "1":
    pytest.skip(
        "set RUN_POSTGRES_INTEGRATION=1 to use the local pgvector database",
        allow_module_level=True,
    )


def _write_representative_files(tmp_path):
    text = tmp_path / "stage1-integration.txt"
    text.write_text("Unique Stage 1 plain text evidence.", encoding="utf-8")

    markdown = tmp_path / "stage1-integration.md"
    markdown.write_text("# Recovery\nUnique recovery evidence.", encoding="utf-8")

    docx = tmp_path / "stage1-integration.docx"
    document = Document()
    document.add_heading("Escalation", level=1)
    document.add_paragraph("Unique DOCX evidence.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Severity"
    table.cell(0, 1).text = "Owner"
    document.save(docx)

    pdf = tmp_path / "stage1-integration.pdf"
    pdf_document = pymupdf.open()
    page = pdf_document.new_page()
    page.insert_text((72, 72), "Unique PDF evidence.")
    pdf_document.save(pdf)
    pdf_document.close()
    return [text, markdown, docx, pdf]


def test_supported_formats_round_trip_through_pgvector(tmp_path):
    created_ids = []

    def fake_embed(*, model, input, truncate):
        return {"embeddings": [[0.01] * 768 for _ in input]}

    try:
        with patch("app.ingestion.pipeline.model_client.embed", side_effect=fake_embed):
            for path in _write_representative_files(tmp_path):
                result = ingest_document(path)
                created_ids.append(result["document_id"])
                assert result["chunk_count"] >= 1
                assert result["created"] is True
    finally:
        for document_id in created_ids:
            delete_document(document_id)
        close_database_pool()
