from docx import Document

from app.ingestion.chunker import chunk_pages
from app.ingestion.parser import parse_docx, parse_markdown, parse_text


def test_text_is_unpaginated(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("A plain text policy.", encoding="utf-8")
    records = parse_text(path)
    assert records[0]["page"] is None
    assert records[0]["content"] == "A plain text policy."


def test_markdown_retains_heading_sections(tmp_path):
    path = tmp_path / "policy.md"
    path.write_text("# Access\nUse SSO.\n## Recovery\nCall support.", encoding="utf-8")
    records = parse_markdown(path)
    assert [record["section"] for record in records] == ["Access", "Recovery"]
    assert "Use SSO." in records[0]["content"]


def test_docx_preserves_paragraph_and_table_order(tmp_path):
    path = tmp_path / "runbook.docx"
    document = Document()
    document.add_heading("Retries", level=1)
    document.add_paragraph("Retry twice.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Code"
    table.cell(0, 1).text = "Action"
    document.save(path)

    records = parse_docx(path)
    assert [record["content"] for record in records] == [
        "Retries", "Retry twice.", "Code\tAction"
    ]
    assert all(record["section"] == "Retries" for record in records)


def test_chunks_keep_nullable_page_and_section():
    chunks = chunk_pages([{
        "filename": "policy.md", "page": None,
        "section": "Access", "content": "Use SSO.",
    }])
    assert chunks[0]["page"] is None
    assert chunks[0]["section"] == "Access"
